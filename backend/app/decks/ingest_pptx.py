"""PPTX ingestion: python-pptx text + notes, LibreOffice -> PDF -> images, text-only fallback."""

import logging
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError

from app.decks.ingest_pdf import IngestError, ingest_pdf
from app.decks.models import Slide

log = logging.getLogger(__name__)

TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
MAX_BULLETS = 8


@dataclass
class PptxSlideText:
    title: str = ""
    bullets: list[str] = field(default_factory=list)
    notes: str = ""


@dataclass
class PptxResult:
    slides: list[Slide]
    rendered: bool
    warning: str | None = None


def find_soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def _shape_lines(shape) -> list[str]:
    lines: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if text:
                lines.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    if hasattr(shape, "shapes"):  # group shape
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
    return lines


def extract_pptx_text(path: Path) -> list[PptxSlideText]:
    """Title placeholder, body text and speaker notes for every visible slide."""
    try:
        prs = Presentation(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise IngestError("Could not read the PPTX file.") from exc

    result: list[PptxSlideText] = []
    for slide in prs.slides:
        if slide._element.get("show") == "0":  # hidden slides are skipped by the PDF export too
            continue
        info = PptxSlideText()
        for shape in slide.shapes:
            is_title = shape.is_placeholder and shape.placeholder_format.type in TITLE_PLACEHOLDERS
            lines = _shape_lines(shape)
            if is_title and lines and not info.title:
                info.title = " ".join(lines)[:120]
            else:
                info.bullets.extend(lines)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            info.notes = slide.notes_slide.notes_text_frame.text.strip()
        result.append(info)
    return result


def convert_to_pdf(path: Path, timeout_s: int = 60) -> tuple[Path | None, str | None]:
    """Run `soffice --headless --convert-to pdf`. Returns (pdf_path, warning)."""
    soffice = find_soffice()
    if not soffice:
        return None, "LibreOffice (soffice) is not installed; showing text-only slides."
    workdir = Path(tempfile.mkdtemp(prefix="pptx2pdf-"))
    try:
        cmd = [
            soffice,
            f"-env:UserInstallation=file://{workdir / 'profile'}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(workdir),
            str(path),
        ]
        subprocess.run(cmd, check=True, capture_output=True, timeout=timeout_s)
        pdf = workdir / (path.stem + ".pdf")
        if not pdf.exists():
            return None, "LibreOffice did not produce a PDF; showing text-only slides."
        target = path.with_suffix(".rendered.pdf")
        shutil.move(str(pdf), target)
        return target, None
    except subprocess.TimeoutExpired:
        return None, "LibreOffice timed out; showing text-only slides."
    except (subprocess.CalledProcessError, OSError) as exc:
        log.warning("soffice failed: %s", exc)
        return None, "LibreOffice failed to render the deck; showing text-only slides."
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def _text_slide(i: int, info: PptxSlideText) -> Slide:
    bullets = [b[:200] for b in info.bullets]
    return Slide(
        index=i,
        title=info.title or (bullets[0][:120] if bullets else f"Slide {i + 1}"),
        bullets=bullets[:MAX_BULLETS],
        text="\n".join(bullets),
        notes=info.notes,
    )


def ingest_pptx(
    pptx_path: Path,
    image_dir: Path,
    deck_id: str,
    *,
    max_slides: int = 40,
    timeout_s: int = 60,
) -> PptxResult:
    texts = extract_pptx_text(pptx_path)
    if not texts:
        raise IngestError("The PPTX has no visible slides.")
    if len(texts) > max_slides:
        raise IngestError(f"The deck has {len(texts)} slides; the limit is {max_slides}.")

    text_slides = [_text_slide(i, info) for i, info in enumerate(texts)]
    pdf, warning = convert_to_pdf(pptx_path, timeout_s=timeout_s)
    if pdf is None:
        return PptxResult(slides=text_slides, rendered=False, warning=warning)

    try:
        rendered = ingest_pdf(pdf, image_dir, deck_id, max_slides=max_slides)
    except IngestError as exc:
        log.warning("Rendered PDF could not be ingested: %s", exc)
        return PptxResult(
            slides=text_slides,
            rendered=False,
            warning="The rendered PDF could not be read; showing text-only slides.",
        )

    merged: list[Slide] = []
    for i, base in enumerate(text_slides):
        page = rendered[i] if i < len(rendered) else None
        merged.append(
            base.model_copy(
                update={
                    "title": texts[i].title or (page.title if page else base.title),
                    "bullets": base.bullets or (page.bullets if page else []),
                    "text": base.text or (page.text if page else ""),
                    "image_url": page.image_url if page else None,
                }
            )
        )
    return PptxResult(slides=merged, rendered=True)
