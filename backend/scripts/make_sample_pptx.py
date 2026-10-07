"""Generate a small sample .pptx (with speaker notes) for upload testing.

Usage (from backend/): uv run python scripts/make_sample_pptx.py [output_path]
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

SLIDES = [
    (
        "Solar Energy 101",
        ["How sunlight becomes electricity", "Why costs fell so fast", "What comes next"],
        "Introduce the topic and the three questions we will answer.",
    ),
    (
        "How Photovoltaics Work",
        ["Photons knock electrons loose in silicon", "Cells wire into panels", "Inverters make AC"],
        "Keep it simple: light in, current out.",
    ),
    (
        "The Cost Curve",
        ["Module prices fell roughly 90% in a decade", "Scale manufacturing", "Better efficiency"],
        "Emphasize learning-rate effects from manufacturing scale.",
    ),
    (
        "Storage and the Grid",
        ["Batteries shift midday power to evening", "Grid upgrades", "Demand flexibility"],
        "Storage is what turns cheap daytime solar into round-the-clock power.",
    ),
    (
        "Summary",
        ["Physics is simple", "Economics now favor solar", "Storage is the next frontier"],
        "Recap the three points and invite questions.",
    ),
]


def build(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, (title, bullets, notes) in enumerate(SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        if i == 0:
            body.text = " · ".join(bullets)
        else:
            body.text = bullets[0]
            for b in bullets[1:]:
                body.add_paragraph().text = b
        slide.notes_slide.notes_text_frame.text = notes
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


if __name__ == "__main__":
    default = Path(__file__).resolve().parents[1] / "samples" / "sample_deck.pptx"
    out = build(Path(sys.argv[1]) if len(sys.argv) > 1 else default)
    print(f"Wrote {out}")
